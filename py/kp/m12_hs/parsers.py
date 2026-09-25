"""M12 table extraction: WCO HS tables, DGFT ITC-HS schedules and correlation tables.

Input is the raw payload of a source file. Excel (.xlsx) and CSV are read; the header row is
located by column-name synonyms within the first rows of each sheet, so column order and extra
columns do not matter. PDF and legacy .xls payloads are rejected with a clear error: the LLD
specifies table extraction from the Excel edition.
"""
from __future__ import annotations

import csv
import io
import re
from dataclasses import dataclass
from typing import Iterable, Iterator, Mapping, Sequence

from .models import (
    CorrelationPair,
    HsCodeRow,
    HsLoadError,
    clean_text,
    level_for,
    normalize_code,
    normalize_policy,
    parent_of,
)

HEADER_SCAN_ROWS = 40
MAX_DESCRIPTION = 4000
MAX_SIMPLE_DESCRIPTION = 1000
MAX_CONDITIONS = 4000

Table = list[list[object]]


# ---- file formats ------------------------------------------------------------------------------

def sniff_format(body: bytes, content_type: str | None = None, url: str | None = None) -> str:
    """'xlsx' | 'csv' | 'pdf' | 'xls'."""
    if body.startswith(b"%PDF"):
        return "pdf"
    if body.startswith(b"PK\x03\x04"):
        return "xlsx"
    if body.startswith(b"\xd0\xcf\x11\xe0"):
        return "xls"
    ct = (content_type or "").lower()
    path = (url or "").lower().split("?")[0]
    if "pdf" in ct or path.endswith(".pdf"):
        return "pdf"
    return "csv"


def read_tables(body: bytes, content_type: str | None = None, url: str | None = None) -> list[Table]:
    """Every sheet of the payload as a list of rows of cell values."""
    fmt = sniff_format(body, content_type, url)
    if fmt == "pdf":
        raise HsLoadError("PDF schedules are not parsed; register the Excel (.xlsx) edition of this table",
                          url=url, format="pdf")
    if fmt == "xls":
        raise HsLoadError("Legacy .xls workbooks are not supported; supply the .xlsx edition", url=url, format="xls")
    if fmt == "xlsx":
        return _read_xlsx(body)
    return [_read_csv(body)]


def _read_xlsx(body: bytes) -> list[Table]:
    from openpyxl import load_workbook

    try:
        wb = load_workbook(io.BytesIO(body), read_only=True, data_only=True)
    except Exception as e:  # noqa: BLE001 — any openpyxl failure means an unreadable workbook
        raise HsLoadError(f"Workbook could not be read: {e}") from e
    try:
        return [[list(r) for r in ws.iter_rows(values_only=True)] for ws in wb.worksheets]
    finally:
        wb.close()


def _read_csv(body: bytes) -> Table:
    for enc in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            text = body.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    else:  # pragma: no cover — latin-1 decodes any byte string
        raise HsLoadError("CSV could not be decoded")
    sample = text[:8192]
    try:
        dialect: type[csv.Dialect] | csv.Dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    return [list(r) for r in csv.reader(io.StringIO(text), dialect)]


# ---- header detection ---------------------------------------------------------------------------

def norm_header(cell: object) -> str:
    return re.sub(r"[^a-z0-9]+", " ", str(cell or "").lower()).strip()


@dataclass(frozen=True)
class ColumnSpec:
    name: str
    exact: frozenset[str]
    contains: tuple[str, ...] = ()
    excludes: tuple[str, ...] = ()
    required: bool = True

    def matches(self, header: str) -> bool:
        if not header:
            return False
        if any(x in header for x in self.excludes):
            return False
        return header in self.exact or any(c in header for c in self.contains)


def locate_columns(table: Table, specs: Sequence[ColumnSpec]) -> tuple[int, dict[str, int]] | None:
    """(header row index, {spec name: column index}) for the first row in which every required
    column is found. Specs are matched in order and a column is used at most once."""
    for i, row in enumerate(table[:HEADER_SCAN_ROWS]):
        headers = [norm_header(c) for c in row]
        used: set[int] = set()
        found: dict[str, int] = {}
        for spec in specs:
            for j, h in enumerate(headers):
                if j not in used and spec.matches(h):
                    found[spec.name] = j
                    used.add(j)
                    break
        if all(s.name in found for s in specs if s.required):
            return i, found
    return None


def _cell(row: Sequence[object], idx: int | None) -> object:
    if idx is None or idx >= len(row):
        return None
    return row[idx]


CODE_EXACT = frozenset({
    "code", "hs code", "hs codes", "hs", "h s", "h s code", "heading", "subheading", "heading subheading",
    "hs heading", "commodity code", "tariff item", "tariff code", "itc hs", "itc hs code", "itchs", "itchs code",
    "itc hs codes", "hs code 8 digit", "cth", "cti",
})
SIMPLE = ColumnSpec("simple", frozenset({"simple description", "plain description"}), ("simple", "plain english"),
                    required=False)
DESCRIPTION = ColumnSpec("description", frozenset({"desc", "goods", "item description", "article description"}),
                         ("description",), excludes=("simple", "plain english"))
CODE = ColumnSpec("code", CODE_EXACT, ("hs code", "itc hs", "itc", "tariff item"), excludes=("policy", "description"))
POLICY = ColumnSpec("policy", frozenset({"policy", "export policy"}), ("export policy", "policy"),
                    excludes=("condition", "notification", "url", "link"), required=False)
CONDITIONS = ColumnSpec("conditions", frozenset({"conditions", "policy conditions"}), ("condition",), required=False)
URL = ColumnSpec("url", frozenset({"url", "link", "source"}), ("url", "link", "notification"), required=False)


# ---- WCO HS nomenclature -----------------------------------------------------------------------

def parse_wco_table(tables: Iterable[Table], version: str) -> list[HsCodeRow]:
    """Chapters, headings and subheadings from a WCO HS table (2/4/6-digit codes; other rows are
    skipped). Raises HsLoadError when no sheet has recognisable code and description columns."""
    out: list[HsCodeRow] = []
    any_header = False
    for table in tables:
        loc = locate_columns(table, [SIMPLE, DESCRIPTION, CODE])
        if loc is None:
            continue
        any_header = True
        hdr, cols = loc
        for row in table[hdr + 1 :]:
            code = normalize_code(_cell(row, cols["code"]))
            if code is None or len(code) == 8:
                continue
            desc = clean_text(_cell(row, cols["description"]), MAX_DESCRIPTION)
            if desc is None:
                continue
            out.append(HsCodeRow(
                version=version, code=code, level=level_for(code), parent_code=parent_of(code), description=desc,
                description_en_simple=clean_text(_cell(row, cols.get("simple")), MAX_SIMPLE_DESCRIPTION),
            ))
    if not any_header:
        raise HsLoadError("No sheet has recognisable HS code and description columns", version=version)
    return out


# ---- DGFT ITC-HS schedule ----------------------------------------------------------------------

def parse_itchs_schedule(tables: Iterable[Table], version: str, source_url: str | None) -> tuple[list[HsCodeRow], list[str]]:
    """8-digit national lines with export policy, conditions and a source link. Rows with other
    code lengths (section, chapter and heading lines in the schedule) are skipped. Returns the rows
    and per-row problems (unrecognised policy wording)."""
    out: list[HsCodeRow] = []
    problems: list[str] = []
    any_header = False
    for table in tables:
        loc = locate_columns(table, [SIMPLE, DESCRIPTION, CODE, CONDITIONS, POLICY, URL])
        if loc is None:
            continue
        any_header = True
        hdr, cols = loc
        for row in table[hdr + 1 :]:
            code = normalize_code(_cell(row, cols["code"]))
            if code is None or len(code) != 8:
                continue
            desc = clean_text(_cell(row, cols["description"]), MAX_DESCRIPTION)
            if desc is None:
                problems.append(f"{code}: no description")
                continue
            try:
                policy = normalize_policy(_cell(row, cols.get("policy")))
            except Exception as e:  # noqa: BLE001 — collected and reported together
                problems.append(f"{code}: {e}")
                continue
            row_url = clean_text(_cell(row, cols.get("url")), 2000)
            url = row_url if row_url and re.match(r"^https?://", row_url) else source_url
            out.append(HsCodeRow(
                version=version, code=code, level="national8", parent_code=code[:6], description=desc,
                description_en_simple=clean_text(_cell(row, cols.get("simple")), MAX_SIMPLE_DESCRIPTION),
                export_policy=policy,
                policy_conditions=clean_text(_cell(row, cols.get("conditions")), MAX_CONDITIONS),
                policy_source_url=url,
            ))
    if not any_header:
        raise HsLoadError("No sheet has recognisable ITC(HS) code and description columns", version=version)
    return out, problems


# ---- correlation tables ------------------------------------------------------------------------

def _year(version: str) -> str:
    m = re.search(r"([0-9]{4})$", version)
    return m.group(1) if m else version


def parse_correlation_table(tables: Iterable[Table], from_version: str, to_version: str) -> list[CorrelationPair]:
    """(from_code, to_code) pairs. Columns are the ones whose header names the version's year
    (e.g. 'HS 2022' / 'HS 2027'). A cell may list several codes ('0101.21; 0101.29')."""
    fy, ty = _year(from_version), _year(to_version)
    if fy == ty:
        # Same year (e.g. HS2022 → ITCHS2022): match the family names instead.
        fy, ty = norm_header(from_version.rstrip("0123456789")), norm_header(to_version.rstrip("0123456789"))
    # 'hs' is a substring of 'itchs'/'itc hs': the HS side must not pick the ITC-HS column.
    from_ex = ("description", "itc") if fy == "hs" else ("description",)
    to_ex = ("description", "itc") if ty == "hs" else ("description",)
    from_spec = ColumnSpec("from", frozenset(), (fy,), excludes=from_ex)
    to_spec = ColumnSpec("to", frozenset(), (ty,), excludes=to_ex)
    pairs: list[CorrelationPair] = []
    any_header = False
    for table in tables:
        loc = locate_columns(table, [from_spec, to_spec])
        if loc is None:
            continue
        any_header = True
        hdr, cols = loc
        last_from: list[str] = []
        for row in table[hdr + 1 :]:
            froms = _codes_in(_cell(row, cols["from"]))
            tos = _codes_in(_cell(row, cols["to"]))
            if froms:
                last_from = froms
            elif tos and last_from:
                # A split listed on continuation rows with the source cell left blank.
                froms = last_from
            for f in froms:
                for t in tos:
                    pairs.append(CorrelationPair(f, t))
    if not any_header:
        raise HsLoadError(f"No sheet has columns for both {from_version} and {to_version}",
                          from_version=from_version, to_version=to_version)
    return pairs


_SPLIT_RE = re.compile(r"[;,\n]+")


def _codes_in(cell: object) -> list[str]:
    if cell is None:
        return []
    if isinstance(cell, (int, float)) and not isinstance(cell, bool):
        c = normalize_code(cell)
        return [c] if c else []
    out: list[str] = []
    for part in _SPLIT_RE.split(str(cell)):
        c = normalize_code(part)
        if c:
            out.append(c)
    return out


def iter_values(rows: Iterable[HsCodeRow]) -> Iterator[Mapping[str, object]]:
    for r in rows:
        yield r.as_value()
